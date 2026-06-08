# extractor/docx_extractor.py
from __future__ import annotations

import logging
from pathlib import Path
from typing import List

from extractor.base import BaseExtractor, register_extractor
from extractor.cleaner import clean_text
from utils.models import DocumentRecord, ExtractionResult

logger = logging.getLogger(__name__)

try:  # pragma: no cover
    from docx import Document as DocxDocument  # type: ignore
except Exception:  # pragma: no cover
    DocxDocument = None  # type: ignore[misc]

try:  # pragma: no cover
    import textract  # type: ignore[attr-defined]
except Exception:  # pragma: no cover
    textract = None  # type: ignore[misc]


class DOCDocxExtractor(BaseExtractor):
    name = "doc"

    DOCX_MEDIA_TYPES = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/x-docx",
    }
    DOC_MEDIA_TYPES = {
        "application/msword",
        "application/doc",
        "application/x-msword",
        "application/vnd.ms-word",
    }
    SUPPORTED_EXTENSIONS = {".docx", ".doc"}

    def supports(self, document: DocumentRecord) -> bool:
        path = document.path
        media_type = (document.media_type or "").lower()

        if path is not None and path.suffix.lower() in self.SUPPORTED_EXTENSIONS:
            return True
        if media_type in self.DOCX_MEDIA_TYPES | self.DOC_MEDIA_TYPES:
            return True
        return False

    def extract(self, document: DocumentRecord) -> ExtractionResult:
        path = self._ensure_path(document)

        if not path.exists():
            raise FileNotFoundError(f"Document path does not exist: {path}")

        ext = path.suffix.lower()
        media_type = (document.media_type or "").lower()
        method: str

        if ext == ".docx" or media_type in self.DOCX_MEDIA_TYPES:
            chunks = self._extract_docx_text(path)
            method = "python-docx"
        elif ext == ".doc" or media_type in self.DOC_MEDIA_TYPES:
            chunks = self._extract_doc_text(path)
            method = "textract"
        else:
            raise ValueError(f"Unsupported Word document format for: {path}")

        raw_text = "\n".join(chunks)
        cleaned_text = clean_text(raw_text)

        if document.metadata is not None:
            document.metadata["extraction_method"] = method
        if ext:
            document.metadata["source_extension"] = ext.lstrip(".")
        else:
            document.metadata.pop("source_extension", None)



        return ExtractionResult(
            document=document,
            raw_text=raw_text,
            clean_text=cleaned_text,
        )

    @staticmethod
    def _ensure_path(document: DocumentRecord) -> Path:
        if document.path is None:
            raise ValueError("DocumentRecord.path must be set for extraction.")
        return document.path

    def _extract_docx_text(self, path: Path) -> List[str]:
        if DocxDocument is None:
            raise RuntimeError(
                "Extracting DOCX files requires the optional 'python-docx' dependency."
            )

        doc = DocxDocument(str(path))
        chunks: List[str] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                chunks.append(text)

        for table in getattr(doc, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    cell_text = "\n".join(
                        p.text.strip()
                        for p in cell.paragraphs
                        if p.text and p.text.strip()
                    )
                    if cell_text:
                        chunks.append(cell_text)

        return chunks

    def _extract_doc_text(self, path: Path) -> List[str]:
        if textract is None:
            raise RuntimeError(
                "Extracting legacy '.doc' files requires the optional 'textract' dependency."
            )

        logger.debug("Using textract for DOC file: %s", path)
        raw_bytes = textract.process(str(path))  # type: ignore[call-arg]
        text = raw_bytes.decode("utf-8", errors="ignore")
        return [text] if text else []


register_extractor("doc", DOCDocxExtractor)
register_extractor("docx", DOCDocxExtractor)
