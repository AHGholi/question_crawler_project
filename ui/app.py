# ui\app.py
"""Streamlit-based UI for running the document question-generation pipeline.

This module wires together the pipeline execution, question rendering, and
export helpers so users can inspect results and download them in common formats.
"""

from __future__ import annotations

import io
import re
import json
import traceback
from typing import Any, Iterable, Optional

import streamlit as st
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

import sys
from pathlib import Path

# Ensure project root is on PYTHONPATH when running `streamlit run ui/app.py`
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from main import run_main_pipeline
from main_pipeline import MainPipelineResult
from processor.pipeline import PipelineContext


# -----------------------------
# Compatibility helpers
# -----------------------------
def _question_text(q: Any) -> str:
    """Return the question text from a question-like object regardless of attribute name."""
    return (getattr(q, "question", None) or getattr(q, "prompt", None) or "").strip()


def _question_answer(q: Any) -> str:
    """Return the answer text from a question-like object if available."""
    return (getattr(q, "answer", None) or "").strip()


def _question_metadata(q: Any) -> dict:
    """Return metadata from a question object when it is stored as a dictionary."""
    md = getattr(q, "metadata", None)
    return md if isinstance(md, dict) else {}


def _iter_questions(ctx: PipelineContext) -> Iterable[Any]:
    """Yield the question items stored in a pipeline context if present."""
    if not ctx or not ctx.questions:
        return []
    items = getattr(ctx.questions, "questions", None)
    return items if items else []


def _safe_title(ctx: PipelineContext) -> str:
    """Return a display-friendly title for the current document context."""
    title = getattr(ctx.document, "title", None) if ctx and ctx.document else None
    if title and str(title).strip():
        return str(title).strip()
    return f"Document {getattr(ctx.document, 'id', 'unknown')}"


# -----------------------------
# UI sanitization helpers
# -----------------------------
_UI_MULTI_SPACE = re.compile(r"\s+")
_UI_LEADING_ENUM = re.compile(
    r"^\s*(?:q(?:uestion)?\s*\d*[:\-\).]\s*|\d+[\).:\-]\s*|[-•]\s*)",
    re.IGNORECASE,
)
_UI_BAD_Q_PAT = re.compile(
    r"(?:\banswer\s*:|\ba\s*:|\bno answer\b|\bnot available\b|\|)",
    re.IGNORECASE,
)

# Split embedded enumerators: Q2. / Q 2 - / Question 3:
_UI_INLINE_Q_SPLIT = re.compile(
    r"(?<!\w)\s*(?:Q|Question)\s*\d+\s*[\).:\-]\s*",
    re.IGNORECASE,
)

# Catch lines where multiple questions are glued by repeated question stems
# e.g. "... Which ... ? Which ... ?"  or "... What ...? What ...?"
_UI_REPEAT_STEM_SPLIT = re.compile(
    r"(?<=[\.\?])\s+(?=(?:What|Which|Why|How|When|Where|Who|Can|Could|Should|Would|Is|Are|Do|Does|Did)\b)",
    re.IGNORECASE,
)


def _clean_question_text(text: str) -> str:
    """Normalize a raw question string before it is shown in the UI."""
    t = (text or "").strip()
    if not t:
        return ""

    t = _UI_LEADING_ENUM.sub("", t).strip()
    t = _UI_MULTI_SPACE.sub(" ", t).strip()

    # remove trailing decorative punctuation
    t = re.sub(r"[.!]+$", "", t).strip()

    # if no question mark but looks like interrogative, add '?'
    if t and not t.endswith("?"):
        if re.match(
            r"^(?:what|which|why|how|when|where|who|whom|whose|is|are|do|does|did|can|could|should|would)\b",
            t,
            flags=re.IGNORECASE,
        ):
            t += "?"
    return t


def _is_displayable_question(text: str) -> bool:
    """Decide whether a cleaned question should be displayed to the user."""
    t = (text or "").strip()
    if not t:
        return False
    if _UI_BAD_Q_PAT.search(t):
        return False

    # Must contain exactly one terminal '?'
    if t.count("?") != 1:
        return False
    if not t.endswith("?"):
        return False

    # reject very short fragments
    if len(t.split()) < 6:
        return False

    # reject dangling endings often produced by truncation
    if re.search(r"\b(with|of|to|for|in|on|at|by|from)\?$", t, flags=re.IGNORECASE):
        return False

    return True


def _clean_answer_text(text: str) -> str:
    """Normalize answer text and suppress placeholder values such as 'N/A'."""
    t = (text or "").strip()
    if not t:
        return ""
    tl = t.lower()
    if tl in {"no answer", "n/a", "na", "unknown", "not available", "none"}:
        return ""
    if re.search(r"^\s*(?:answer|a)\s*:\s*$", t, flags=re.IGNORECASE):
        return ""
    return _UI_MULTI_SPACE.sub(" ", t).strip()


def _split_merged_questions(text: str) -> list[str]:
    """Split merged question text that contains multiple questions in one block."""
    t = (text or "").strip()
    if not t:
        return []

    t = _UI_MULTI_SPACE.sub(" ", t).strip()

    # Normalize hard separators first
    t = t.replace("||", " | ").replace(" | ", "\n")

    # 1) split at explicit Q markers
    t = _UI_INLINE_Q_SPLIT.sub("\n", t)

    # 2) split after '?' if next token starts a new likely question stem
    t = re.sub(
        r"(?<=\?)\s+(?=(?:What|Which|Why|How|When|Where|Who|Can|Could|Should|Would|Is|Are|Do|Does|Did)\b)",
        "\n",
        t,
        flags=re.IGNORECASE,
    )

    # 3) split repeated stems after punctuation
    t = _UI_REPEAT_STEM_SPLIT.sub("\n", t)

    parts = [p.strip(" -•\t\r\n") for p in t.split("\n") if p.strip()]
    out: list[str] = []

    for p in parts:
        # If multiple '?' remain, keep up to first full question only
        if p.count("?") > 1:
            p = p.split("?", 1)[0].strip() + "?"
        out.append(p)

    return out


def _visible_questions(ctx: PipelineContext) -> list[tuple[str, str, dict]]:
    """Return the cleaned questions that should be rendered for a pipeline context."""
    out: list[tuple[str, str, dict]] = []

    for q in _iter_questions(ctx):
        raw = _question_text(q)
        md = _question_metadata(q)
        ans = _clean_answer_text(_question_answer(q))

        # break merged candidates
        pieces = _split_merged_questions(raw) or [raw]

        for piece in pieces:
            qt = _clean_question_text(piece)
            if not _is_displayable_question(qt):
                continue
            out.append((qt, ans, md))

    # UI-layer dedupe (case-insensitive exact)
    seen = set()
    deduped: list[tuple[str, str, dict]] = []
    for qt, ans, md in out:
        k = qt.lower()
        if k in seen:
            continue
        seen.add(k)
        deduped.append((qt, ans, md))
    return deduped


# -----------------------------
# Export helpers
# -----------------------------
def question_set_to_docx(contexts: list[PipelineContext]) -> bytes:
    """Export the visible questions for multiple contexts into a DOCX document."""
    doc = DocxDocument()
    doc.add_heading("Generated Questions", level=1)

    total = 0
    section_num = 0
    for ctx in contexts:
        visible = _visible_questions(ctx)
        if not visible:
            continue

        section_num += 1
        doc.add_heading(f"{section_num}. {_safe_title(ctx)}", level=2)
        for j, (text, _ans, md) in enumerate(visible, start=1):
            doc.add_paragraph(f"Q{j}. {text}")
            if md:
                doc.add_paragraph(f"Metadata: {json.dumps(md, ensure_ascii=False)}")
            total += 1

    if total == 0:
        doc.add_paragraph("No questions generated.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def question_set_to_pdf(contexts: list[PipelineContext]) -> bytes:
    """Export the visible questions for multiple contexts into a PDF document."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    page_w, page_h = A4

    left_margin = 40
    right_margin = 40
    top_margin = 40
    bottom_margin = 40

    font_name = "Helvetica"
    font_size = 11
    line_h = 16
    y = page_h - top_margin
    max_width = page_w - left_margin - right_margin

    try:
        pdfmetrics.registerFont(TTFont("DejaVu", "DejaVuSans.ttf"))
        font_name = "DejaVu"
    except Exception:
        font_name = "Helvetica"

    c.setFont(font_name, font_size)

    def _new_page() -> None:
        nonlocal y
        c.showPage()
        c.setFont(font_name, font_size)
        y = page_h - top_margin

    def _wrap_text(text: str, width: float) -> list[str]:
        """Word-wrap text to fit PDF width; also hard-wrap very long tokens."""
        text = (text or "").replace("\t", "    ")
        raw_lines = text.splitlines() or [""]
        out: list[str] = []

        for raw in raw_lines:
            words = raw.split(" ")
            cur = ""
            for w in words:
                candidate = w if not cur else f"{cur} {w}"
                if pdfmetrics.stringWidth(candidate, font_name, font_size) <= width:
                    cur = candidate
                else:
                    if cur:
                        out.append(cur)
                        cur = ""

                    # hard-wrap single overlong token (e.g. long URL / JSON)
                    if pdfmetrics.stringWidth(w, font_name, font_size) > width:
                        chunk = ""
                        for ch in w:
                            test = chunk + ch
                            if pdfmetrics.stringWidth(test, font_name, font_size) <= width:
                                chunk = test
                            else:
                                if chunk:
                                    out.append(chunk)
                                chunk = ch
                        cur = chunk
                    else:
                        cur = w
            out.append(cur)
        return out

    def write_line(text: str = "") -> None:
        nonlocal y
        wrapped = _wrap_text(text, max_width) if text else [""]
        for line in wrapped:
            if y < bottom_margin:
                _new_page()
            c.drawString(left_margin, y, line)
            y -= line_h

    write_line("Generated Questions")
    write_line("-" * 40)

    total = 0
    section_num = 0
    for ctx in contexts:
        visible = _visible_questions(ctx)
        if not visible:
            continue

        section_num += 1
        write_line(f"{section_num}. {_safe_title(ctx)}")

        for j, (text, _ans, md) in enumerate(visible, start=1):
            write_line(f"  Q{j}. {text}")
            if md:
                write_line(f"     Metadata: {json.dumps(md, ensure_ascii=False)}")
            total += 1

        write_line("")

    if total == 0:
        write_line("No questions generated.")

    c.save()
    buf.seek(0)
    return buf.getvalue()



# -----------------------------
# UI Rendering
# -----------------------------
def render_context(ctx: PipelineContext, idx: int) -> None:
    """Render the question results for one document context in the Streamlit UI."""
    title = _safe_title(ctx)
    with st.expander(f"{idx}. {title}", expanded=False):
        if ctx.errors:
            st.error("Errors:")
            for err in ctx.errors:
                st.code(err)

        extraction = ctx.extraction
        if extraction:
            clean_len = len((getattr(extraction, "clean_text", "") or "").strip())
            kws = getattr(extraction, "keywords", None) or []
            top_sents = getattr(extraction, "top_sentences", None) or []
            st.caption(f"clean_text_len={clean_len} | keywords={len(kws)} | top_sentences={len(top_sents)}")

        visible = _visible_questions(ctx)
        if not visible:
            st.info("No questions generated for this document.")
            return

        for q_idx, (q_text, _q_ans, q_md) in enumerate(visible, start=1):
            st.markdown(f"**Q{q_idx}. {q_text}**")
            if q_md:
                with st.expander("Metadata", expanded=False):
                    st.json(q_md)


def _stats_block(result: MainPipelineResult) -> None:
    """Render a small summary block for the pipeline execution statistics."""
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Total documents", result.stats.total_documents)
    col2.metric("Processed", result.stats.processed_documents)
    col3.metric("Failed", result.stats.failed_documents)
    col4.metric("Total questions", result.stats.total_questions)


# -----------------------------
# Main app
# -----------------------------
def main() -> None:
    """Run the Streamlit application and handle the full user workflow."""
    st.set_page_config(page_title="Question Generator", layout="wide")
    st.title("Document Question Generator")

    with st.sidebar:
        st.header("Run settings")
        query = st.text_input("Query", value="machine learning basics")
        max_results = st.number_input("Max search results", min_value=1, max_value=50, value=5, step=1)
        limit = st.number_input("Process limit (0 = no limit)", min_value=0, max_value=100, value=0, step=1)

        st.divider()
        st.header("Question generation backend")

        qg_provider = st.selectbox(
            "Backend",
            options=["local_hf", "hf_api"],
            format_func=lambda x: "Local Hugging Face" if x == "local_hf" else "Hugging Face API",
            index=0,
        )

        default_model = "google/flan-t5-large"
        qg_model = st.text_input("Model", value=default_model)

        with st.expander("Advanced generation settings", expanded=False):
            qg_temperature = st.slider("Temperature", min_value=0.0, max_value=1.5, value=0.7, step=0.1)
            qg_max_new_tokens = st.number_input("Max new tokens", min_value=16, max_value=1024, value=320, step=8)

            if qg_provider == "local_hf":
                qg_min_new_tokens = st.number_input("Min new tokens", min_value=1, max_value=256, value=12, step=1)
                qg_do_sample = st.checkbox("Do sample", value=True)
                qg_top_p = st.slider("Top-p", min_value=0.1, max_value=1.0, value=0.9, step=0.05)
                qg_num_beams = st.number_input("Beam count", min_value=1, max_value=16, value=4, step=1)
                qg_device = st.number_input("Device (-1 = CPU, 0 = first GPU)", min_value=-1, max_value=8, value=-1, step=1)
                qg_timeout = None
            else:
                qg_min_new_tokens = None
                qg_do_sample = None
                qg_top_p = None
                qg_num_beams = None
                qg_device = None
                qg_timeout = st.number_input("API timeout (seconds)", min_value=5, max_value=300, value=60, step=5)

        if qg_provider == "hf_api":
            qg_api_token = st.text_input("HF API token", value="", type="password")
            st.caption("If left blank, the app will try the HF_API_TOKEN value from your .env/environment.")
        else:
            qg_api_token = None

        if "is_generating" not in st.session_state:
            st.session_state["is_generating"] = False

        run_btn = st.button("Run pipeline", type="primary", disabled=st.session_state["is_generating"])

    if "last_result" not in st.session_state:
        st.session_state["last_result"] = None
    if "last_error" not in st.session_state:
        st.session_state["last_error"] = None
    if "exports" not in st.session_state:
        st.session_state["exports"] = None

    if run_btn and not st.session_state["is_generating"]:
        st.session_state["is_generating"] = True
        st.session_state["last_error"] = None
        with st.spinner("Running pipeline..."):
            try:
                limit_arg: Optional[int] = None if int(limit) == 0 else int(limit)

                result = run_main_pipeline(
                    query=query.strip(),
                    max_results=int(max_results),
                    limit=limit_arg,
                    qg_provider=qg_provider,
                    qg_model=qg_model.strip(),
                    qg_api_token=(qg_api_token.strip() if qg_api_token else None),
                    qg_temperature=float(qg_temperature),
                    qg_max_new_tokens=int(qg_max_new_tokens),
                    qg_min_new_tokens=(int(qg_min_new_tokens) if qg_min_new_tokens is not None else None),
                    qg_do_sample=qg_do_sample,
                    qg_top_p=(float(qg_top_p) if qg_top_p is not None else None),
                    qg_num_beams=(int(qg_num_beams) if qg_num_beams is not None else None),
                    qg_device=(int(qg_device) if qg_device is not None else None),
                    qg_timeout=(int(qg_timeout) if qg_timeout is not None else None),
                )
                st.session_state["last_result"] = result

                contexts_with_questions = [c for c in result.contexts if len(_visible_questions(c)) > 0]

                raw = []
                for ctx in contexts_with_questions:
                    visible = _visible_questions(ctx)
                    raw.append(
                        {
                            "document_id": getattr(ctx.document, "id", None),
                            "title": _safe_title(ctx),
                            "questions": [
                                {
                                    "question": text,
                                    "answer": "",
                                    "metadata": md,
                                }
                                for (text, _ans, md) in visible
                            ],
                        }
                    )

                st.session_state["exports"] = {
                    "has_questions": len(contexts_with_questions) > 0,
                    "docx": question_set_to_docx(contexts_with_questions),
                    "pdf": question_set_to_pdf(contexts_with_questions),
                    "json": json.dumps(raw, ensure_ascii=False, indent=2).encode("utf-8"),
                }

            except Exception as exc:
                st.session_state["last_result"] = None
                st.session_state["exports"] = None
                st.session_state["last_error"] = f"{exc}\n\n{traceback.format_exc()}"
            finally:
                st.session_state["is_generating"] = False

    if st.session_state["last_error"]:
        st.error("Pipeline failed")
        st.code(st.session_state["last_error"])
        return

    result: MainPipelineResult | None = st.session_state["last_result"]
    if not result:
        st.info("Set your query and click **Run pipeline**.")
        return

    _stats_block(result)

    st.divider()
    st.subheader("Per-document results")
    if not result.contexts:
        st.warning("No documents returned.")
    else:
        for i, ctx in enumerate(result.contexts, start=1):
            render_context(ctx, i)

    st.divider()
    st.subheader("Export")

    exports = st.session_state.get("exports") or {
        "has_questions": False,
        "docx": b"",
        "pdf": b"",
        "json": b"",
    }

    c1, c2, c3 = st.columns([1, 1, 2])

    with c1:
        st.download_button(
            "Download DOCX",
            data=exports["docx"],
            file_name="generated_questions.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            disabled=not exports["has_questions"],
            key="download_docx",
        )

    with c2:
        st.download_button(
            "Download PDF",
            data=exports["pdf"],
            file_name="generated_questions.pdf",
            mime="application/pdf",
            disabled=not exports["has_questions"],
            key="download_pdf",
        )

    with c3:
        st.download_button(
            "Download JSON",
            data=exports["json"],
            file_name="generated_questions.json",
            mime="application/json",
            disabled=not exports["has_questions"],
            key="download_json",
        )


if __name__ == "__main__":
    main()
